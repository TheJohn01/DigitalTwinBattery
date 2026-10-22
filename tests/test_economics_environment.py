import numpy as np
import pandas as pd

from battery_twin.chemistry import CHEMISTRIES
from battery_twin.economics import (FIXED_ELECTRICITY_PRICE, FIXED_SPREAD, GRID_FEE_SCENARIOS,
                                    NREL_SYSTEM_RTE, POWER_CONVERSION_RTE, TAX_LOSS_CAP,
                                    arbitrage_prices, best_retirement, calculate_tea,
                                    daily_arbitrage, irc_rate, second_life_vs_new, sensitivity,
                                    system_capex_eur_kwh, tea_at_best_retirement,
                                    tea_documentation)
from battery_twin.environment import CFB_USE_KEY, calculate_lca, total_energy_kwh
from battery_twin.lca_report import STATEMENT, build_lca_report, lca_results
from battery_twin.tea_report import build_tea_report

NEW_FOREVER = np.ones(20000)                    # a battery that never ages
# Settings of one dashboard run and LCA inputs, shared by the report tests
RUN = dict(c_charge=0.3, c_discharge=1.0, ambient_c=20, cycling_c=20,
           cooling="Forced air / passive", h=25, mean_cell_temp=22, cell_capacity_ah=4.6,
           cell_energy_wh=16.9, cycles_per_year=365, model="SPMe", n_cycles=1000,
           cycles_to_80=374)
LCA_ARGS = dict(eol_soh=0.7, service_years=15, cycles_per_year=365, round_trip_eff=0.9,
                grid_g_co2_per_kwh=210, lifetime_km=200000, kwh_per_km=0.16)


def linear_fade(per_cycle):
    """Capacity curve losing per_cycle each cycle, cut at 50 % like the dashboard."""
    q = 1 - per_cycle * np.arange(1, 20001)
    return np.where(q >= 0.5, q, np.nan)


def test_npv_falls_when_capex_rises():
    assert (calculate_tea(NEW_FOREVER, capex_eur_kwh=200)["NPV [€]"]
            < calculate_tea(NEW_FOREVER, capex_eur_kwh=100)["NPV [€]"])


def test_battery_without_ageing_is_never_retired():
    assert calculate_tea(NEW_FOREVER)["Retired in year"] is None


def test_worn_out_battery_has_no_costs_after_retirement():
    # 0.1 %/cycle at 365 cycles/year: 50 % is reached in year 2 (cycle 501). Afterwards the
    # cash flow must be zero, not "operating costs of a dead battery".
    r = calculate_tea(linear_fade(0.001), years=15, retire_below=0.5)
    assert r["Retired in year"] == 2
    r_short = calculate_tea(linear_fade(0.001), years=3, retire_below=0.5)
    assert abs(r["NPV [€]"] - r_short["NPV [€]"]) < 1e-6


def test_slower_fade_gives_higher_npv():
    assert (calculate_tea(linear_fade(0.0001))["NPV [€]"]
            > calculate_tea(linear_fade(0.001))["NPV [€]"])


def test_second_life_starts_at_the_right_capacity():
    sl, _ = second_life_vs_new({"soh_curve": linear_fade(0.0001)}, linear_fade(0.0001)[1000:],
                               20.0, linear_fade(0.0001), 2034, 8)
    assert abs(sl["SOH at start [%]"] - 89.99) < 1e-9


def test_second_life_and_new_system_each_retire_at_their_own_best_threshold():
    # One end-of-life rule for both: with grid fees the best threshold is not the end of the
    # data, and it differs between the two curves.
    fees = GRID_FEE_SCENARIOS["BTE access tariff from 2030"]
    base = {"soh_curve": linear_fade(0.0004), "spread": 0.06, "retire_below": 0.6,
            "grid_fee_eur_kwh": fees[0], "grid_fee_eur_kw_day": fees[1]}
    used_curve, new_curve = linear_fade(0.0004)[2000:], linear_fade(0.0001)
    sl, new = second_life_vs_new(base, used_curve, 150.0, new_curve, 2034, 8)
    common = dict(base, start_year=2034, years=8)
    _, sl_best, _, _ = tea_at_best_retirement(dict(common, soh_curve=used_curve,
                                                   capex_eur_kwh=150.0))
    _, new_best, _, _ = tea_at_best_retirement(dict(common, soh_curve=new_curve,
                                                    capex_eur_kwh=system_capex_eur_kwh("LFP"),
                                                    round_trip_eff=NREL_SYSTEM_RTE))
    assert np.isclose(sl["NPV [€]"], sl_best["NPV [€]"])
    assert np.isclose(new["NPV [€]"], new_best["NPV [€]"])


def test_repurposed_battery_cost_is_nrel_b2u():
    from battery_twin.economics import repurposed_pack_usd_kwh
    assert abs(repurposed_pack_usd_kwh(0.70) - 44.0) < 0.01      # B2U forced selling price


def test_lca_carbon_footprint_positive():
    for chem in CHEMISTRIES:
        _, ind = calculate_lca(chem, 60)
        assert ind[CFB_USE_KEY] > 0


def test_npv_is_always_a_number():
    # Even a curve that ends immediately (all NaN) must give a finite NPV.
    for curve in [linear_fade(0.001), linear_fade(0.01), np.full(20000, np.nan)]:
        assert np.isfinite(calculate_tea(curve)["NPV [€]"])


def test_total_energy_follows_cfb_rules():
    q = np.linspace(1.0, 0.5, 1001)               # loses 0.05 % per cycle
    energy, n = total_energy_kwh(10, q, eol_soh=0.7)
    assert n == 601                                # cycle 602 is the first below 70 %
    assert abs(energy - 10 * q[:601].sum()) < 1e-9
    # service-life cap: 1 year at 365 cycles/year
    assert total_energy_kwh(10, q, 0.7, service_years=1, cycles_per_year=365)[1] == 365


def test_lca_uses_degen_data_and_reports_water_and_ecotox():
    phases, ind = calculate_lca("LFP", 10, throughput_kwh=10000)
    prod = phases.iloc[0]
    assert abs(prod["Emissions [kg CO2eq]"] - 10 * 58.5) < 1e-6
    assert abs(prod["Water consumption [m3]"] - 10 * 0.91) < 1e-6
    assert abs(prod["Freshwater ecotoxicity [kg 1,4-DCB eq]"] - 10 * 16.0) < 1e-6
    assert ind["Water consumption [L/kWh delivered]"] > 0


def test_lca_report_pdf():
    for chem in ["NMC", "LFP", "NCA", "Na-ion"]:
        pdf = build_lca_report(chem, 60, np.linspace(1, 0.5, 5000), "stationary", RUN, LCA_ARGS)
        assert pdf.startswith(b"%PDF") and len(pdf) > 5000
    assert "not critically reviewed" in STATEMENT


def _prices(day_prices, days=3):
    """Hourly price table with the same 24 prices [€/MWh] every day."""
    return pd.DataFrame([{"date": f"2025-01-{d:02d}", "hour": h + 1, "price_eur_mwh": p}
                         for d in range(1, days + 1) for h, p in enumerate(day_prices)])


def test_arbitrage_buys_cheap_hours_and_sells_dear_hours():
    day = [50.0] * 24
    day[3] = day[4] = 10.0          # cheap at night
    day[19] = 200.0                 # one dear evening hour
    daily = daily_arbitrage(_prices(day), c_charge=0.5, c_discharge=1.0, dod=1.0, rte=1.0)
    assert len(daily) == 3
    assert abs(daily["buy [€/kWh]"].iloc[0] - 0.010) < 1e-9      # 2 h at €10/MWh
    assert abs(daily["sell [€/kWh]"].iloc[0] - 0.200) < 1e-9     # 1 h at €200/MWh
    arb = arbitrage_prices(daily, 365)
    assert abs(arb["spread"] - 0.190) < 1e-9 and arb["profitable_days"] == 3


def test_arbitrage_charges_before_it_discharges():
    # Dear in the morning, cheap in the evening: the battery cannot sell before buying.
    day = [200.0] * 12 + [10.0] * 12
    daily = daily_arbitrage(_prices(day, 1), c_charge=1.0, c_discharge=1.0, dod=1.0, rte=1.0)
    assert daily["margin [€/kWh]"].iloc[0] <= 1e-9


def test_sourced_costs():
    # NREL $321/kWh system + Benchmark cell price difference to LFP, at 1.13 $/€
    assert abs(system_capex_eur_kwh("LFP") - 321 / 1.13) < 1e-6
    assert abs(system_capex_eur_kwh("NMC") - (321 + 66.0 - 50.2) / 1.13) < 1e-6
    assert abs(system_capex_eur_kwh("NCA") - (321 + 71.7 - 50.2) / 1.13) < 1e-6
    assert system_capex_eur_kwh("Na-ion") == system_capex_eur_kwh("LFP")
    assert abs(POWER_CONVERSION_RTE - 83 / 91) < 1e-12
    assert calculate_tea(NEW_FOREVER)["CAPEX [€]"] == system_capex_eur_kwh("LFP") * 60


def test_grid_fees_only_after_the_2029_exemption():
    energy, power = GRID_FEE_SCENARIOS["MT access tariff from 2030"]
    t = calculate_tea(NEW_FOREVER, grid_fee_eur_kwh=energy, grid_fee_eur_kw_day=power)["Table"]
    assert (t.loc[t["Calendar year"] <= 2029, "Grid fees [€]"] == 0).all()
    assert (t.loc[t["Calendar year"] >= 2030, "Grid fees [€]"] > 0).all()
    exempt = calculate_tea(NEW_FOREVER)["Table"]
    assert (exempt["Grid fees [€]"] == 0).all()


def test_income_tax_rates_and_losses_carried_forward():
    assert (irc_rate(2026), irc_rate(2027), irc_rate(2030)) == (0.19, 0.18, 0.17)
    profitable = calculate_tea(NEW_FOREVER, spread=0.3)
    t = profitable["Table"]
    assert (t["Income tax [€]"] >= 0).all() and t["Income tax [€]"].sum() > 0
    no_tax = calculate_tea(NEW_FOREVER, spread=0.3, income_tax=False)
    assert no_tax["NPV [€]"] > profitable["NPV [€]"]
    loss_making = calculate_tea(NEW_FOREVER, spread=0.0)["Table"]
    assert loss_making["Income tax [€]"].sum() == 0


def test_cash_flow_table_adds_up():
    r = calculate_tea(linear_fade(0.0002), spread=0.1)
    t = r["Table"]
    ebitda = t["Revenue [€]"] - t["Charging cost [€]"] - t["O&M and warranty [€]"] - t["Grid fees [€]"]
    assert np.allclose(ebitda, t["EBITDA [€]"])
    assert np.allclose(t["EBITDA [€]"] - t["Income tax [€]"], t["Cash flow [€]"])
    assert abs(r["Cash flow [€]"][0] + r["CAPEX [€]"]) < 1e-9
    doc = tea_documentation({"soh_curve": None, "capex_eur_kwh": 284.1}, "LFP")
    assert "ISO/TS 14076" in doc["method"] and doc["assumptions"]


def test_tax_depreciation_and_loss_cap():
    r = calculate_tea(NEW_FOREVER, spread=0.3, years=10)
    dep = r["Table"]["Tax depreciation [€]"]
    assert np.allclose(dep[:5], 0.2 * r["CAPEX [€]"]) and np.allclose(dep[5:], 0)   # 20 %/year
    slow = calculate_tea(NEW_FOREVER, spread=0.3, years=10, depreciation_rate=0.1)
    assert np.allclose(slow["Table"]["Tax depreciation [€]"], 0.1 * r["CAPEX [€]"])
    # A loss in year 1 can offset at most 65 % of the taxable income of a later year.
    t = calculate_tea(NEW_FOREVER, spread=0.3, years=10, depreciation_rate=1.0)["Table"]
    taxable_2 = t["EBITDA [€]"].iloc[1]
    assert t["Income tax [€]"].iloc[1] >= 0.18 * (1 - TAX_LOSS_CAP) * taxable_2 - 1e-6


def test_tea_report_pdf_and_word():
    q = linear_fade(0.0002)
    inputs = dict(soh_curve=q, capex_eur_kwh=284.1, capacity_kwh=60, electricity_price=0.026,
                  spread=0.095, round_trip_eff=0.85, cycles_per_year=365)
    lca = lca_results("LFP", 60, q, "stationary")["ind"]
    for fee in GRID_FEE_SCENARIOS:
        pdf = build_tea_report("LFP", RUN, inputs, None, fee, lca)
        assert pdf[:4] == b"%PDF"
    word = build_tea_report("LFP", RUN, inputs, None, fee, lca, fmt="docx")
    assert word[:2] == b"PK"                       # a .docx is a zip file
    assert build_lca_report("LFP", 60, q, "stationary", RUN, LCA_ARGS, fmt="docx")[:2] == b"PK"


def test_best_retirement_is_the_highest_npv():
    inputs = {"soh_curve": linear_fade(0.0001), "spread": 0.05,
              "grid_fee_eur_kwh": GRID_FEE_SCENARIOS["MT access tariff from 2030"][0],
              "grid_fee_eur_kw_day": GRID_FEE_SCENARIOS["MT access tariff from 2030"][1]}
    table, best = best_retirement(inputs)
    assert best["NPV [€]"] == table["NPV [€]"].max()
    check = calculate_tea(**dict(inputs, retire_below=best["Retire below [%]"] / 100))
    assert np.isclose(check["NPV [€]"], best["NPV [€]"])


def test_without_running_costs_keep_the_battery_to_the_end_of_the_data():
    _, best = best_retirement({"soh_curve": linear_fade(0.0001), "om_eur_kwh_yr": 0.0,
                               "warranty_pct": 0.0})
    assert best["Retire below [%]"] == 50


# --- Regressions for the redundancy review (October 2026) ---------------------------------
def test_sensitivity_keeps_a_zero_input():
    # A1: a spread or O&M of 0 is an input, not a missing value. ±20 % of 0 is 0, so the
    # NPV must not change.
    for key in ("spread", "om_eur_kwh_yr"):
        s = sensitivity({"soh_curve": linear_fade(0.0002), key: 0.0}).set_index("Parameter")
        row = s.loc["Arbitrage spread" if key == "spread" else "O&M cost"]
        assert abs(row["ΔNPV at -20% [€]"]) < 1e-6 and abs(row["ΔNPV at +20% [€]"]) < 1e-6


def test_sensitivity_uses_the_defaults_of_calculate_tea():
    # Leaving an input out must give the same tornado as passing its default explicitly.
    q = linear_fade(0.0002)
    implicit = sensitivity({"soh_curve": q})
    explicit = sensitivity({"soh_curve": q, "spread": FIXED_SPREAD,
                            "electricity_price": FIXED_ELECTRICITY_PRICE,
                            "capex_eur_kwh": system_capex_eur_kwh("LFP")})
    assert np.allclose(implicit.iloc[:, 1:], explicit.iloc[:, 1:])


def test_one_definition_of_cycles_to_a_level():
    # A2: 80 % is reached exactly at cycle 400, so the first cycle BELOW 80 % is 401, and
    # the LCA counts the 400 cycles before it.
    from battery_twin.ageing import cycles_to
    q = 1 - 0.0005 * np.arange(1, 2001)
    assert cycles_to(q, 0.8) == 401
    assert total_energy_kwh(1, q, 0.8)[1] == 400
    assert cycles_to(np.full(10, np.nan), 0.8) == 1            # no data: below from the start
    assert cycles_to(np.ones(10), 0.8) == float("inf")


def test_tea_report_retirement_cycle_is_one_based():
    # A2: the report's "Cycles to 70 % (retirement)" uses the same definition as "Cycles to
    # 80 %": first cycle below. 0.0005 per cycle reaches 70 % exactly at cycle 600.
    from docx import Document
    from io import BytesIO
    q = np.where(1 - 0.0005 * np.arange(1, 20001) >= 0.5, 1 - 0.0005 * np.arange(1, 20001),
                 np.nan)
    inputs = dict(soh_curve=q, capacity_kwh=60, cycles_per_year=365, round_trip_eff=0.85)
    word = build_tea_report("LFP", RUN, inputs, None, list(GRID_FEE_SCENARIOS)[0], None,
                            fmt="docx")
    cells = [c.text for t in Document(BytesIO(word)).tables for r in t.rows for c in r.cells]
    assert cells[cells.index("Cycles to 70 % (retirement)") + 1] == "601"


def test_tea_report_reuses_given_results():
    # B1: the dashboard passes its TEA results; the report must not recompute them.
    import battery_twin.tea_report as tea_report
    q = linear_fade(0.0002)
    inputs = dict(soh_curve=q, capacity_kwh=60, cycles_per_year=365, round_trip_eff=0.85)
    tea = calculate_tea(**inputs)
    sens = sensitivity(inputs, base_npv=tea["NPV [€]"])
    retirement, _ = best_retirement(inputs)
    calls = []
    original = tea_report.calculate_tea
    tea_report.calculate_tea = lambda **kw: calls.append(1) or original(**kw)
    try:
        build_tea_report("LFP", RUN, inputs, None, list(GRID_FEE_SCENARIOS)[0], None,
                         tea=tea, sens=sens, retirement=retirement)
    finally:
        tea_report.calculate_tea = original
    assert len(calls) == 3                     # only the 3 scenarios that differ from the base


def test_lca_results_follow_their_inputs():
    # A3: the grid intensity chosen by the user must reach the LCA (and so the TEA report).
    q = linear_fade(0.0002)
    low = lca_results("NMC", 60, q, "stationary", **dict(LCA_ARGS, grid_g_co2_per_kwh=210))
    high = lca_results("NMC", 60, q, "stationary", **dict(LCA_ARGS, grid_g_co2_per_kwh=400))
    assert high["ind"][CFB_USE_KEY] > low["ind"][CFB_USE_KEY]


# --- One end of life (highest NPV) and the efficiency each report uses --------------------
def _docx_text(data):
    from io import BytesIO
    from docx import Document
    doc = Document(BytesIO(data))
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return " ".join(p.text for p in doc.paragraphs), cells


def test_tea_at_best_retirement_uses_the_highest_npv():
    inputs = {"soh_curve": linear_fade(0.0001), "spread": 0.05,
              "grid_fee_eur_kwh": GRID_FEE_SCENARIOS["MT access tariff from 2030"][0],
              "grid_fee_eur_kw_day": GRID_FEE_SCENARIOS["MT access tariff from 2030"][1]}
    used, tea, table, best = tea_at_best_retirement(inputs)
    assert used["retire_below"] == best["Retire below [%]"] / 100
    assert np.isclose(tea["NPV [€]"], table["NPV [€]"].max())


def test_tea_report_states_the_end_of_life_and_its_efficiency():
    inputs = dict(soh_curve=linear_fade(0.0002), capacity_kwh=60, cycles_per_year=365,
                  round_trip_eff=0.83)
    used, tea, table, best = tea_at_best_retirement(inputs)
    text, cells = _docx_text(build_tea_report("LFP", dict(RUN, rte_cell=0.91), used, None,
                                              list(GRID_FEE_SCENARIOS)[0], None, fmt="docx",
                                              tea=tea, retirement=table))
    assert "the threshold with the highest NPV" in text
    assert "Round-trip efficiency used: the whole system, 83.0 %" in text
    assert "grid connection" in text and "PNNL-27237" in text
    used_row = cells.index("used")
    assert cells[used_row - 4] == f"{best['Retire below [%]']:.0f} %"


def test_lca_report_states_the_end_of_life_and_its_efficiency():
    run = dict(RUN, rte_cell=0.91, rte_system=0.83,
               eol_basis="Retirement threshold with the highest NPV (TEA of this run)")
    text, cells = _docx_text(build_lca_report("NMC", 60, linear_fade(0.0002), "stationary", run,
                                              dict(LCA_ARGS, eol_soh=0.6, round_trip_eff=0.91),
                                              fmt="docx"))
    assert "Round-trip efficiency used: the battery itself (its cells), 91.0 %" in text
    assert "whole-system efficiency (83.0 %)" in text
    assert "End of service life: 60 %" in text and "not a conforming declaration" in text
    assert any("CFB-IND minimum, 70 %" in c for c in cells)          # sensitivity case
    assert any("whole system (83.0 %" in c for c in cells)            # sensitivity case


def test_tea_report_second_life_section():
    second = {"ev_years": 8, "km_day": 50.0, "capacity_left": 0.889, "euro7_ok": True,
              "buy_usd_kwh_left": 28.03, "repurposing_usd_kwh": 24.38, "pack_usd_kwh": 49.3,
              "capex": 266.0, "years": 8, "start_year": 2034, "npv": -8007.0,
              "npv_new": -8480.0, "co2_t": 3.5, "new_capex": 284.0, "eol": 50.0, "eol_new": 55.0,
              "sensitivity": pd.DataFrame({"Case": ["Base case"], "Second life minus new LFP [€]": [473]}),
              "carbon": {"Car life only": 125.0, "Second life alone (cut-off)": 0.1}}
    text, cells = _docx_text(build_tea_report("LFP", RUN, dict(soh_curve=linear_fade(0.0002)),
                                              None, list(GRID_FEE_SCENARIOS)[0], None,
                                              fmt="docx", second=second))
    assert "3.6 Second life" in text and "€473 better" in text.replace(",", "")
    assert "-€8,007" in cells



def test_second_life_carbon_shares_production():
    from battery_twin.environment import second_life_carbon
    c = second_life_carbon("NMC", 60, 23360, 60000, 80000)
    # one battery, more energy delivered: the production per kWh falls; repurposing is tiny
    assert c["Car and second life (production shared)"] < c["Car life only"]
    assert c["Second life alone (cut-off)"] < 1
    lca_docx = build_lca_report("NMC", 60, linear_fade(0.0002), "stationary", RUN, LCA_ARGS,
                                fmt="docx", second={"ev_years": 8, "capacity_left": 0.9,
                                                    "years": 8, "carbon": c})
    assert "4.6 Second life" in _docx_text(lca_docx)[0]


def test_price_index_follows_the_calendar_year():
    # Banco de Portugal projections for 2026-28, then 2 %: a second life starting in 2034
    # must see 2034 prices, not 2025 prices raised by one year
    from battery_twin.economics import price_index_from_base
    assert abs(price_index_from_base(2026) - 1.031) < 1e-12
    assert abs(price_index_from_base(2028) - 1.031 * 1.024 * 1.020) < 1e-12
    assert abs(price_index_from_base(2034) / price_index_from_base(2028) - 1.02 ** 6) < 1e-12


def test_a_later_start_pays_that_years_prices():
    # CAPEX and O&M are 2025 prices; bought in 2034 they cost 2033's price level, like the income
    from battery_twin.economics import calculate_tea, price_index_from_base
    curve = np.full(3000, 0.95)
    now = calculate_tea(curve, start_year=2026, years=3)
    later = calculate_tea(curve, start_year=2034, years=3)
    level = price_index_from_base(2033)
    assert abs(later["CAPEX [€]"] - level * now["CAPEX [€]"]) < 1e-6
    assert abs(later["Table"]["O&M and warranty [€]"].iloc[0]
               - level * now["Table"]["O&M and warranty [€]"].iloc[0]) < 1e-6
